"""Genera los entregables del modelo de datos de FARMIX a partir del catálogo real de PostgreSQL.

Entradas:  catalogo.json (hechos, extraído con catalogo.sql) y descripciones.json (interpretación funcional).
Salidas:   relaciones.csv, diccionario-datos.{csv,md,xlsx}, modelo-fisico.md, modelo-logico*.mmd,
           areas/*.mmd y farmix-drawdb.sql.
Uso:       python3 generar.py
"""
import csv
import json
import re
from pathlib import Path

AQUI = Path(__file__).parent
CAT = json.loads((AQUI / 'catalogo.json').read_text())['tablas']
DES = json.loads((AQUI / 'descripciones.json').read_text())
REGLA = {'a': 'NO ACTION', 'r': 'RESTRICT', 'c': 'CASCADE', 'n': 'SET NULL', 'd': 'SET DEFAULT'}
AUDITORIA = {'created_by', 'updated_by'}
TABLAS = {t['tabla']: t for t in CAT}
AREA_DE = {tb: clave for clave, a in DES['areas'].items() for tb in a['tablas']}


def restr(t, tipo):
    return [k for k in (t['restricciones'] or []) if k['tipo'] == tipo]


def fks(t):
    return restr(t, 'f')


def pk_cols(t):
    p = restr(t, 'p')
    return p[0]['columnas'] if p else []


def fk_de_columna(t, col):
    return [k for k in fks(t) if col in k['columnas']]


def checks_de_columna(t, col):
    return [k['definicion'].replace('CHECK ', '') for k in restr(t, 'c') if re.search(rf'\b{col}\b', k['definicion'])]


def unicos_de_columna(t, col):
    """Restricciones UNIQUE e índices únicos (incluidos los parciales) que contienen la columna."""
    nombres = {k['nombre'] for k in restr(t, 'u')} | {k['nombre'] for k in restr(t, 'p')}
    out = [f"{k['nombre']} ({', '.join(k['columnas'])})" for k in restr(t, 'u') if col in k['columnas']]
    for i in t['indices'] or []:
        # Solo las columnas indexadas (antes de WHERE): la condición de un índice parcial no hace única a la columna
        indexadas = i['definicion'].split(' ON ')[1].split(' WHERE ')[0]
        if 'UNIQUE' in i['definicion'] and i['nombre'] not in nombres and re.search(rf'\b{col}\b', indexadas):
            out.append(i['nombre'] + ' (índice único)')
    return out


def descripcion(tabla, col):
    d = DES['tablas'].get(tabla, {}).get('columnas', {})
    return d.get(col) or DES['genericas'].get(col, '')


def tipo_corto(tipo):
    return tipo.replace('timestamp with time zone', 'timestamptz')


# ---------- relaciones.csv ----------
with open(AQUI / 'relaciones.csv', 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f, delimiter=';')
    w.writerow(['tabla_origen', 'columnas_fk', 'tabla_destino', 'columnas_referenciadas', 'on_delete', 'on_update',
                'obligatoria', 'cardinalidad', 'restriccion', 'tipo'])
    for t in CAT:
        for k in fks(t):
            nulos = any(c['nulo'] for c in t['columnas'] if c['columna'] in k['columnas'])
            w.writerow([t['tabla'], ', '.join(k['columnas']), k['destino'], ', '.join(k['columnas_destino']),
                        REGLA[k['al_borrar']], REGLA[k['al_actualizar']], 'no' if nulos else 'sí',
                        f"N:{'0..1' if nulos else '1'}", k['nombre'],
                        'auditoría' if set(k['columnas']) <= AUDITORIA else ('compuesta' if len(k['columnas']) > 1 else 'simple')])

# ---------- diccionario de datos ----------
filas = []
for t in CAT:
    pk = pk_cols(t)
    for c in t['columnas']:
        col = c['columna']
        claves = (['PK'] if col in pk else []) + (['FK'] if fk_de_columna(t, col) else []) + (['UNIQUE'] if unicos_de_columna(t, col) else [])
        ref = '; '.join(f"{k['destino']}({', '.join(k['columnas_destino'])}) [{', '.join(k['columnas'])} → {REGLA[k['al_borrar']]}]" for k in fk_de_columna(t, col))
        valid = checks_de_columna(t, col) + [f'UNIQUE: {u}' for u in unicos_de_columna(t, col)]
        filas.append({
            'Área': AREA_DE.get(t['tabla'], ''), 'Tabla': t['tabla'], 'Columna': col, 'Tipo de dato': c['tipo'],
            'Clave': ', '.join(claves) or '—', 'Admite NULL': 'sí' if c['nulo'] else 'no',
            'Valor predeterminado': c['defecto'] or ('IDENTITY' if c['identidad'] else ''),
            'Referencia': ref, 'Función (interpretación)': descripcion(t['tabla'], col),
            'Validaciones (catálogo)': ' | '.join(valid), 'Observaciones': ''})

campos = list(filas[0].keys())
with open(AQUI / 'diccionario-datos.csv', 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.DictWriter(f, fieldnames=campos, delimiter=';')
    w.writeheader()
    w.writerows(filas)

md = ['# Diccionario de datos de FARMIX', '',
      'Generado con `generar.py` desde el catálogo de PostgreSQL (13 migraciones aplicadas).',
      'Las columnas **Tipo, Clave, NULL, Predeterminado, Referencia y Validaciones** son hechos del catálogo;',
      '**Función** es interpretación redactada a partir del código y debe leerse como tal.', '']
for t in CAT:
    a = DES['areas'][AREA_DE[t['tabla']]]['nombre']
    md += [f"## `{t['tabla']}`", '', f"**Área:** {a}. **Qué es:** {DES['tablas'][t['tabla']]['que']} **Módulos:** {DES['tablas'][t['tabla']]['modulos']}.", '',
           '| Columna | Tipo | Clave | NULL | Predeterminado | Referencia | Función | Validaciones |', '|---|---|---|---|---|---|---|---|']
    for r in [x for x in filas if x['Tabla'] == t['tabla']]:
        esc = lambda s: str(s).replace('|', '\\|')
        md.append(f"| `{r['Columna']}` | {r['Tipo de dato']} | {r['Clave']} | {r['Admite NULL']} | {esc(r['Valor predeterminado'])} | "
                  f"{esc(r['Referencia'])} | {esc(r['Función (interpretación)'])} | {esc(r['Validaciones (catálogo)'])} |")
    md.append('')
(AQUI / 'diccionario-datos.md').write_text('\n'.join(md), encoding='utf-8')

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Diccionario'
    ws.append(campos)
    for r in filas:
        ws.append([r[c] for c in campos])
    for celda in ws[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='1F7A5C')
    anchos = [8, 22, 24, 26, 14, 10, 30, 40, 60, 60, 25]
    for i, ancho in enumerate(anchos):
        ws.column_dimensions[chr(65 + i)].width = ancho
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = ws.dimensions
    rel = wb.create_sheet('Relaciones')
    with open(AQUI / 'relaciones.csv', encoding='utf-8-sig') as f:
        for fila in csv.reader(f, delimiter=';'):
            rel.append(fila)
    for celda in rel[1]:
        celda.font = Font(bold=True)
    wb.save(AQUI / 'diccionario-datos.xlsx')
except ImportError:
    print('openpyxl no está instalado: se omite el Excel')

# ---------- modelo físico ----------
fis = ['# Modelo físico de FARMIX en PostgreSQL', '',
       'Generado desde el catálogo. Muestra cómo se implementa cada tabla: tipos, valores predeterminados,',
       'restricciones e índices exactamente como los tiene PostgreSQL.', '']
for t in CAT:
    fis += [f"## `{t['tabla']}`", '', '| Columna | Tipo | NULL | Predeterminado |', '|---|---|---|---|']
    for c in t['columnas']:
        fis.append(f"| `{c['columna']}` | {c['tipo']} | {'sí' if c['nulo'] else 'no'} | {(c['defecto'] or ('IDENTITY' if c['identidad'] else '')).replace('|', '\\|')} |")
    fis += ['', '**Restricciones**', '']
    nombre_tipo = {'p': 'PRIMARY KEY', 'f': 'FOREIGN KEY', 'u': 'UNIQUE', 'c': 'CHECK'}
    for k in t['restricciones'] or []:
        fis.append(f"- `{k['nombre']}` ({nombre_tipo[k['tipo']]}): `{k['definicion']}`")
    fis += ['', '**Índices**', '']
    for i in t['indices'] or []:
        fis.append(f"- `{i['definicion'].replace('public.', '')}`")
    fis.append('')
(AQUI / 'modelo-fisico.md').write_text('\n'.join(fis), encoding='utf-8')


# ---------- Mermaid (modelo lógico) ----------
def entidad_mermaid(t, solo_claves=False):
    pk, lineas = pk_cols(t), []
    for c in t['columnas']:
        col = c['columna']
        marcas = (['PK'] if col in pk else []) + (['FK'] if fk_de_columna(t, col) else []) + (['UK'] if unicos_de_columna(t, col) else [])
        if solo_claves and not marcas:
            continue
        tipo = re.sub(r'\(.*\)', '', tipo_corto(c['tipo'])).replace(' ', '_')
        nota = '' if c['nulo'] is False else ' "opcional"'
        lineas.append(f"    {tipo} {col} {','.join(marcas)}{nota}".rstrip())
    return [f"  {t['tabla']} {{"] + lineas + ['  }']


def relacion_mermaid(t, k):
    nulos = any(c['nulo'] for c in t['columnas'] if c['columna'] in k['columnas'])
    padre = '|o' if nulos else '||'
    return f"  {k['destino']} {padre}--o{{ {t['tabla']} : \"{', '.join(k['columnas'])}\""


def diagrama(tablas, con_auditoria, solo_claves=False, externos=()):
    out = ['erDiagram']
    for nombre in tablas:
        out += entidad_mermaid(TABLAS[nombre], solo_claves)
    for nombre in externos:
        out += [f"  {nombre} {{", f"    uuid id PK \"ver otra área\"", '  }']
    visibles = set(tablas) | set(externos)
    for nombre in tablas:
        for k in fks(TABLAS[nombre]):
            if (con_auditoria or not set(k['columnas']) <= AUDITORIA) and k['destino'] in visibles:
                out.append(relacion_mermaid(TABLAS[nombre], k))
    return '\n'.join(out) + '\n'


todas = [t['tabla'] for t in CAT]
(AQUI / 'modelo-logico-completo.mmd').write_text(diagrama(todas, True), encoding='utf-8')
(AQUI / 'modelo-logico-general.mmd').write_text(diagrama(todas, False, solo_claves=True), encoding='utf-8')
(AQUI / 'areas').mkdir(exist_ok=True)
for clave, a in DES['areas'].items():
    externos = sorted({k['destino'] for tb in a['tablas'] for k in fks(TABLAS[tb])
                       if k['destino'] not in a['tablas'] and not set(k['columnas']) <= AUDITORIA})
    (AQUI / 'areas' / f"{clave}.mmd").write_text(diagrama(a['tablas'], False, externos=externos), encoding='utf-8')

# ---------- SQL para drawDB ----------
# Solo CREATE TABLE con columnas, PK, UNIQUE y FK, en un formato sencillo para importadores de diagramas.
# No es una migración: omite CHECK, índices parciales o con expresiones y valores predeterminados con funciones.
draw = ['-- FARMIX: esquema simplificado para dibujar el diagrama en drawDB (NO usar como migración).',
        '-- Generado por generar.py. Omite: restricciones CHECK, índices (incluidos los únicos parciales y con lower()),',
        '-- valores predeterminados con funciones (gen_random_uuid(), now()) e IDENTITY. Ver drawdb-diferencias.md.', '']
for t in CAT:
    cols = []
    for c in t['columnas']:
        tipo = tipo_corto(c['tipo']).replace('timestamptz', 'timestamp')
        defecto = c['defecto'] if c['defecto'] and '(' not in c['defecto'] else None
        cols.append(f"  {c['columna']} {tipo}{'' if c['nulo'] else ' NOT NULL'}{f' DEFAULT {defecto}' if defecto else ''}")
    cols.append(f"  PRIMARY KEY ({', '.join(pk_cols(t))})")
    for u in restr(t, 'u'):
        cols.append(f"  UNIQUE ({', '.join(u['columnas'])})")
    draw += [f"CREATE TABLE {t['tabla']} (", ',\n'.join(cols), ');', '']
for t in CAT:
    for k in fks(t):
        draw.append(f"ALTER TABLE {t['tabla']} ADD FOREIGN KEY ({', '.join(k['columnas'])}) REFERENCES {k['destino']} "
                    f"({', '.join(k['columnas_destino'])}) ON DELETE {REGLA[k['al_borrar']]};")
(AQUI / 'farmix-drawdb.sql').write_text('\n'.join(draw) + '\n', encoding='utf-8')

print(f'{len(CAT)} tablas, {len(filas)} columnas, {sum(len(fks(t)) for t in CAT)} relaciones')
